"""
Export Service — Excel (openpyxl) and PDF (reportlab) generation for reports
"""
import io
from decimal import Decimal
from datetime import date, datetime
from typing import List, Dict, Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle


# ── Excel Export ───────────────────────────────────────────────

def export_to_excel(title: str, columns: List[str], column_labels: List[str], rows: List[Dict[str, Any]]) -> bytes:
    """Export generic tabular data to styled Excel file, returns bytes"""
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"

    header_font = Font(name="Tahoma", bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill(start_color="00D4AA", end_color="00D4AA", fill_type="solid")
    title_font  = Font(name="Tahoma", bold=True, size=14)
    normal_font = Font(name="Tahoma", size=10)
    border = Border(bottom=Side(style="thin", color="DDDDDD"))

    # Title row
    ws.merge_cells(f"A1:{get_column_letter(len(columns))}1")
    ws["A1"] = title
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(horizontal="center")

    # Generated date
    ws.merge_cells(f"A2:{get_column_letter(len(columns))}2")
    ws["A2"] = f"สร้างเมื่อ: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    ws["A2"].font = Font(name="Tahoma", size=9, italic=True, color="888888")
    ws["A2"].alignment = Alignment(horizontal="center")

    # Header row
    header_row = 4
    for i, label in enumerate(column_labels, start=1):
        cell = ws.cell(row=header_row, column=i, value=label)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Data rows
    for r, row_data in enumerate(rows, start=header_row + 1):
        for c, col in enumerate(columns, start=1):
            value = row_data.get(col)
            cell = ws.cell(row=r, column=c, value=value)
            cell.font = normal_font
            cell.border = border
            if isinstance(value, (int, float)):
                cell.number_format = "#,##0.00"
                cell.alignment = Alignment(horizontal="right")

    # Auto-width columns
    for i, label in enumerate(column_labels, start=1):
        max_len = max([len(str(label))] + [len(str(r.get(columns[i-1], ""))) for r in rows[:100]])
        ws.column_dimensions[get_column_letter(i)].width = min(max_len + 4, 40)

    ws.freeze_panes = f"A{header_row + 1}"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def export_pl_to_excel(pl_data: dict) -> bytes:
    """Export P&L statement to formatted Excel"""
    wb = Workbook()
    ws = wb.active
    ws.title = "งบกำไรขาดทุน"

    title_font  = Font(name="Tahoma", bold=True, size=14)
    section_font = Font(name="Tahoma", bold=True, size=11, color="FFFFFF")
    section_fill = PatternFill(start_color="151C28", end_color="151C28", fill_type="solid")
    total_font  = Font(name="Tahoma", bold=True, size=11)
    normal_font = Font(name="Tahoma", size=10)

    ws.merge_cells("A1:C1")
    ws["A1"] = "งบกำไรขาดทุน (Profit & Loss Statement)"
    ws["A1"].font = title_font
    ws.merge_cells("A2:C2")
    ws["A2"] = f"{pl_data['period_from']} ถึง {pl_data['period_to']}"
    ws["A2"].font = Font(name="Tahoma", size=10, italic=True)

    row = 4

    def write_section(section, row):
        ws.cell(row=row, column=1, value=section["label"]).font = section_font
        ws.cell(row=row, column=1).fill = section_fill
        ws.merge_cells(f"A{row}:C{row}")
        row += 1
        for item in section["items"]:
            ws.cell(row=row, column=1, value=item["account_code"]).font = normal_font
            ws.cell(row=row, column=2, value=item["account_name"]).font = normal_font
            c = ws.cell(row=row, column=3, value=item["amount"])
            c.font = normal_font
            c.number_format = "#,##0.00"
            row += 1
        ws.cell(row=row, column=2, value=f"รวม{section['label']}").font = total_font
        c = ws.cell(row=row, column=3, value=section["total"])
        c.font = total_font
        c.number_format = "#,##0.00"
        row += 2
        return row

    row = write_section(pl_data["revenue"], row)
    row = write_section(pl_data["cogs"], row)

    ws.cell(row=row, column=2, value="กำไรขั้นต้น").font = total_font
    c = ws.cell(row=row, column=3, value=pl_data["gross_profit"])
    c.font = total_font
    c.number_format = "#,##0.00"
    row += 2

    row = write_section(pl_data["expenses"], row)

    ws.cell(row=row, column=2, value="กำไร(ขาดทุน)สุทธิ").font = Font(name="Tahoma", bold=True, size=12)
    c = ws.cell(row=row, column=3, value=pl_data["net_profit"])
    c.font = Font(name="Tahoma", bold=True, size=12)
    c.number_format = "#,##0.00"

    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 32
    ws.column_dimensions["C"].width = 18

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


# ── PDF Export ─────────────────────────────────────────────────

def export_to_pdf(title: str, subtitle: str, columns: List[str], column_labels: List[str], rows: List[Dict[str, Any]]) -> bytes:
    """Export generic tabular data to PDF"""
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=1.5*cm, bottomMargin=1.5*cm)
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle("TitleTH", parent=styles["Title"], fontSize=16, spaceAfter=4)
    sub_style   = ParagraphStyle("SubTH", parent=styles["Normal"], fontSize=9, textColor=colors.grey, spaceAfter=16)

    elements = [Paragraph(title, title_style), Paragraph(subtitle, sub_style)]

    table_data = [column_labels]
    for row in rows:
        table_row = []
        for col in columns:
            val = row.get(col, "")
            if isinstance(val, (int, float)):
                table_row.append(f"{val:,.2f}")
            else:
                table_row.append(str(val) if val is not None else "")
        table_data.append(table_row)

    table = Table(table_data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#00D4AA")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#DDDDDD")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F5F5")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))

    elements.append(table)
    elements.append(Spacer(1, 20))
    elements.append(Paragraph(f"สร้างเมื่อ {datetime.now().strftime('%d/%m/%Y %H:%M')} โดย MeAccount", sub_style))

    doc.build(elements)
    buf.seek(0)
    return buf.read()
